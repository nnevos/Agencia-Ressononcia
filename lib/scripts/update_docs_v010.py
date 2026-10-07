from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path('/mnt/data/v010_refactor')
GDD_IN=Path('/mnt/data/RESSONANCIA_GDD_V3_3_v0_9_13.docx')
AU_IN=Path('/mnt/data/RESSONANCIA_BIBLIA_AU_V2_1_v0_8_1.docx')
GDD_OUT=ROOT/'docs/RESSONANCIA_GDD_V4_0_v0_1_0.docx'
AU_OUT=ROOT/'docs/RESSONANCIA_BIBLIA_AU_V3_0_v0_1_0.docx'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn(f'w:{m}'))
        if node is None: node=OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def style_table(table, widths=None):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.style='Table Grid'
    for r_idx,row in enumerate(table.rows):
        for c_idx,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 90, 100, 90, 100)
            if r_idx==0:
                shade(cell,'102833')
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(9)
            else:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size=Pt(9)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)

def add_heading(doc,text,level=1):
    p=doc.add_heading(text, level=level)
    return p

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet')
        p.add_run(item)

def replace_all(doc, replacements):
    for p in doc.paragraphs:
        full=''.join(r.text for r in p.runs)
        new=full
        for a,b in replacements.items(): new=new.replace(a,b)
        if new!=full:
            for r in p.runs: r.text=''
            if p.runs: p.runs[0].text=new
            else: p.add_run(new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full=''.join(r.text for r in p.runs); new=full
                    for a,b in replacements.items(): new=new.replace(a,b)
                    if new!=full:
                        for r in p.runs: r.text=''
                        if p.runs: p.runs[0].text=new
                        else: p.add_run(new)

def update_gdd():
    doc=Document(GDD_IN)
    replace_all(doc, {
        'Versão 3.3 / projeto v0.9.13':'Versão 4.0 / projeto v0.1.0 Foundation',
        'Pós-expediente Conversas, VN/webnovel, amizade/romance e consequências':'Pós-expediente NEXO privado, cenas presenciais especiais, amizade/romance e consequências',
        ' /conversa — pós-expediente VN':' /conversa — hub pós-expediente NEXO; conversas opcionais com múltiplos personagens',
        'A v0.9.13 usa save schema v6.':'O baseline v0.1.0 mantém save schema v6.',
    })
    doc.add_page_break()
    add_heading(doc,'20. Baseline v0.1.0 — arquitetura orientada a conteúdo',1)
    doc.add_paragraph('A linha funcional v0.9.14 foi consolidada como um novo baseline v0.1.0 Foundation. A mudança de numeração não representa regressão de funcionalidades; representa uma reorganização para transformar o protótipo em uma base sustentável de produção. O antigo snapshot 0.1.1 permanece obsoleto e não deve ser usado como base.')
    t=doc.add_table(rows=1, cols=3)
    for i,v in enumerate(['Camada','Responsabilidade','Regra de autoria']): t.rows[0].cells[i].text=v
    rows=[
        ('content/','Personagens, ocorrências, falas, respostas, tutorial, introdução, textos de UI e balanceamento.','Fonte oficial editável pelo autor.'),
        ('game/','Simulação, turno, condição, resolução, progressão e regras.','Não acumular texto narrativo específico.'),
        ('components/','Componentes reutilizáveis de interface.','Recebem dados; não definem cânone ou conteúdo.'),
        ('app/','Telas e roteamento.','Consomem content/ + game/.'),
        ('lib/','Save e infraestrutura.','Mudanças exigem migração, nunca apagar save silenciosamente.'),
        ('docs/','GDD, AU, autoria e failsafe.','Toda decisão durável deve ser documentada.'),
    ]
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v
    style_table(t,[1.05,2.5,2.8])

    add_heading(doc,'20.1 Controle editorial',2)
    add_bullets(doc,[
        'Ocorrências e seus tempos/requisitos vivem em content/incidents/.',
        'Mensagens do NEXO durante o trabalho vivem em content/messages/operations.ts e podem ser ligadas a eventos ou ser conversas ambientais independentes.',
        'Conversas privadas pós-expediente vivem em content/dialogues/post-shift/, separadas por personagem.',
        'O texto de cada opção do jogador e a resposta do personagem são escritos explicitamente pelo autor; o motor apenas aplica flags e deltas relacionais.',
        'Introdução e tutorial vivem em content/narrative/. Textos permanentes do menu vivem em content/ui/.',
        'Balanceamento central (Vida/Energia, XP, custos e limites de chance) vive em content/config/.',
    ])

    add_heading(doc,'21. Pós-expediente e comunicação social',1)
    doc.add_paragraph('NEXO é a interface social cotidiana. Durante o expediente, o Analista acompanha o canal operacional em modo somente leitura. Fora do expediente, o jogo abre um hub de DMs privadas.')
    add_bullets(doc,[
        'O jogador pode conversar com qualquer personagem que possua cena disponível naquela noite.',
        'Pode conversar com vários personagens na mesma noite.',
        'Pode encerrar a noite sem conversar com ninguém.',
        'Concluir uma DM retorna ao hub; não força automaticamente o avanço do dia.',
        'Disponibilidade de cena pode depender de dia, flags, missões, relação e outros requisitos.',
        'Cenas presenciais/VN continuam como momentos especiais: encontros, conflitos, convites, romance e marcos de campanha.',
    ])

    add_heading(doc,'22. Introdução, tutorial e menu',1)
    doc.add_paragraph('O baseline inclui estrutura data-driven para menu, introdução e tutorial. A introdução usa “Coordenação” como placeholder editorial enquanto a identidade da chefia da Agência permanece [A DEFINIR]. Quando esse personagem for aprovado na AU, a troca deve ocorrer no conteúdo, não dentro da lógica da tela.')
    add_bullets(doc,[
        'Menu principal deve comportar Continuar, Novo jogo e futuras opções normais de configuração/save.',
        'Tutorial deve explicar mapa, briefing, agentes, Vida/Energia, chance estimada e NEXO sem exigir alteração de motor para reescrever texto.',
        'Novas telas narrativas devem preferir dados estruturados e IDs estáveis.',
    ])

    add_heading(doc,'23. Princípio de manutenção',1)
    p=doc.add_paragraph(); r=p.add_run('“Se uma mudança altera o que o jogador lê, encontra ou escolhe, ela deve ser editável em content/. Se altera como o jogo calcula, pertence ao motor.”'); r.bold=True
    doc.add_paragraph('A validação de desenvolvimento verifica erros editoriais básicos, como IDs duplicados e cenas apontando para personagens inexistentes. O TypeScript continua sendo a primeira barreira contra conteúdo estruturalmente inválido.')
    doc.save(GDD_OUT)


def update_au():
    doc=Document(AU_IN)
    replace_all(doc, {'Versão 2.1 / projeto v0.8.1':'Versão 3.0 / projeto v0.1.0 Foundation'})
    doc.add_page_break()
    add_heading(doc,'16. Comunicação e vida social na AU',1)
    doc.add_paragraph('[AU-APROVADO] NEXO é a principal interface cotidiana de comunicação da Agência. Durante o expediente, o Analista acompanha o grupo operacional em modo somente leitura; sua agência acontece principalmente através do despacho. No pós-expediente, o mesmo ecossistema abre conversas privadas nas quais o jogador pode escolher respostas.')
    add_bullets(doc,[
        '[AU-APROVADO] Conversas privadas são opcionais. O Analista pode falar com um, vários ou nenhum personagem numa noite.',
        '[AU-APROVADO] Relações dos sete entre si continuam acontecendo no grupo operacional independentemente da participação do Analista.',
        '[AU-APROVADO] Cenas presenciais/VN são reservadas para momentos que justifiquem sair do celular: encontros, conflitos, convites, romance, investigação e marcos de campanha.',
        '[A DEFINIR] Identidade, nome, aparência e histórico da chefia/coordenação que apresenta formalmente o trabalho ao Analista.',
    ])

    add_heading(doc,'17. Regra de autoria narrativa',1)
    doc.add_paragraph('[AU-APROVADO] A fala dos personagens e as opções do jogador são conteúdo autoral explícito. O motor não deve inventar automaticamente falas, preferências românticas ou fatos de cânone. Cada cena deve registrar texto, requisitos, flags e efeitos de relação.')
    add_bullets(doc,[
        'Falas operacionais: content/messages/.',
        'DMs pós-expediente: content/dialogues/post-shift/.',
        'Introdução/tutorial: content/narrative/.',
        'Antes de escrever cena: CÂNONE-BASE → AU-APROVADO → A DEFINIR.',
        'Uma decisão narrativa durável exige atualização desta Bíblia/AU_CHANGES e do failsafe.',
    ])

    add_heading(doc,'18. Estrutura de campanha editável',1)
    doc.add_paragraph('O baseline v0.1.0 separa conteúdo e motor para permitir que a campanha cresça por arquivos de autoria. Ocorrências, mensagens e cenas devem possuir IDs estáveis. Requisitos de cena podem usar dia e flags; futuramente podem incluir relação, resultado de ocorrência, condição e outros gatilhos sem reescrever a identidade do personagem.')
    doc.save(AU_OUT)

update_gdd(); update_au()
print(GDD_OUT); print(AU_OUT)
